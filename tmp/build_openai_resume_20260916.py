from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('/Users/bmysoren/prep/resume/Bharath_Mysore_OpenAI_Compute_Storage_Resume.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.65)
sec.left_margin = sec.right_margin = Inches(0.7)
sec.footer_distance = Inches(0.3)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(5)
    style.paragraph_format.line_spacing = 1.04
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(6)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)

def p(text, style=None):
    para = doc.add_paragraph(text, style)
    para.paragraph_format.widow_control = True
    return para

def bullet(text):
    para = p(text, 'List Bullet')
    para.paragraph_format.left_indent = Inches(0.15)
    para.paragraph_format.first_line_indent = Inches(-0.15)
    para.paragraph_format.keep_together = True

def role(company, dates, scope):
    para = p('', 'Heading 2')
    para.add_run(company).bold = True
    para.add_run(' | ' + dates).bold = False
    para = p(scope)
    para.paragraph_format.keep_with_next = True
    para.runs[0].italic = True

p('BHARATH MYSORE', 'Title')
p('Bothell, WA | +1 (425) 503-6839 | bharathkumarmn@hotmail.com')
para = p('STORAGE INFRASTRUCTURE AND DISTRIBUTED SYSTEMS')
para.runs[0].bold = True
p('Systems software engineer and technical leader with 20+ years of experience across cloud storage, operating systems, and developer tools. At Oracle Cloud Infrastructure, design and implement storage capabilities spanning object, block, and file workloads, shared replication, cross-region disaster recovery, and capacity efficiency. Combine hands-on C/C++ development with cross-team delivery, deep systems debugging, and automation of storage operations.')
p('Professional Experience', 'Heading 1')
role('Oracle Corporation', 'Jun 2016 - Present', 'Oracle Cloud Infrastructure | Unified Storage and Block Storage')
for text in [
    'Design and drive cross-team initiatives for a unified storage platform serving object, block, and file workloads, including a shared replication layer and erasure coding for object storage.',
    'Build block-storage support on the unified stack and drive migration from the existing backend; contribute to deep archival capabilities for long-term storage.',
    'Develop cross-region replication for disaster recovery, including asynchronous replication of volume groups. Ownership spans replication and leader election, snapshots, backup, and cross-region replication.',
    'Implement application-controlled snapshotting and scale the backup backend with service growth, supporting data protection across the storage lifecycle.',
    'Increase sellable capacity through cold-data compression, reclaiming 20-30% of storage space. Optimize idle-volume backend capacity through work associated with U.S. Patent 11,412,043.',
    'Implement reclamation of trimmed user blocks to recover unreferenced storage capacity.',
    'Scale iSCSI serving through multiprocessing and rearchitect the networking subsystem to reduce hops and increase the number of attached volumes.',
    'Refactor the storage backend to contain failures within a feature, volume, or device. Automate failed-disk handling and detection and handling of single-host failures.',
    'Drive automated qualification and deployment through CI/CD, together with single-click region buildout to reduce manual operational work.',
]:
    bullet(text)

doc.add_page_break()
p('Professional Experience Continued', 'Heading 1')
role('Microsoft Corporation', 'Sep 2005 - May 2016', 'Azure Storage | Visual Studio | Windows Systems | Device Platforms')
for text in [
    'Designed and implemented an RDMA library for a proof-of-concept Azure Storage backend, connecting systems programming experience with storage transport performance.',
    'Led Visual Studio Cluster Debugger for HPC/MPI, Parallel Debugger, and GPU Debugger projects spanning CUDA, OpenGL, and DirectX. Built teams of five local and ten remote engineers; contributed to hiring and mentoring.',
    'Implemented engine test libraries for cluster and GPU debuggers and architected a reusable debugger UI test library. Coordinated with vendors to migrate tests across product cycles.',
    'Diagnosed Windows operating-system, filesystem, disk, and cluster failures through kernel-dump analysis, reverse engineering, and code review. Worked with development teams and customers to resolve root causes.',
    'Owned end-to-end HoloLens depth-camera functionality from prototype through public announcement. Developed hardware-validation drivers and tools for Surface Hub and Kinect, and worked on Xbox graphics emulation and frame-capture tooling.',
]:
    bullet(text)
role('Intel Technology India Private Limited', 'Mar 2002 - Sep 2005', 'Systems Software and Hardware Validation')
for text in [
    'Designed and developed CPU-core and cache/memory-subsystem validation tools, including Linux kernel programming and drivers for test cards and chipsets.',
    'Investigated hardware and operating-system defects, drove issues to closure, and wrote targeted patches. Developed network-card management and configuration software for Windows and Linux.',
]:
    bullet(text)
p('Additional Cloud Service Delivery', 'Heading 1')
p('At Oracle, designed and developed a near-real-time marketplace software-metering service from prototype to public release. Helped redirect and implement a data-processing service, taking it from prototype through internal release.')
p('Technical Skills', 'Heading 1')
for label, value in [
    ('Languages', 'C/C++; Java, Python, and C# as needed.'),
    ('Storage and systems', 'Object, block, and file storage; replication; leader election; erasure coding; snapshots; backup; disaster recovery; compression; iSCSI; RDMA; Linux kernel and device drivers; concurrent and parallel programming.'),
    ('Infrastructure and tooling', 'Kubernetes, Terraform, Docker, Helm, Prometheus, gRPC, Folly, Boost, ZooKeeper, Kafka, and Flink.'),
]:
    para = p('')
    para.add_run(label + ': ').bold = True
    para.add_run(value)
footer = sec.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('Bharath Mysore | ')
run.font.size = Pt(9)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Bharath Mysore Resume for OpenAI Compute Storage'
doc.core_properties.author = 'Bharath Mysore'
doc.core_properties.subject = 'Storage infrastructure and distributed systems'
doc.save(OUT)
print(OUT)
