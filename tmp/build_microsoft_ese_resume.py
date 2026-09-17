from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/bmysoren/prep/resume/Bharath_Mysore_Microsoft_ESE_Storage_Engine.docx"


def set_font(run, name="Calibri", size=10.0, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_spacing(paragraph, before=0, after=3, line=1.05):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_bottom_border(paragraph, color="D9E2F3", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_spacing(p, before=5, after=3, line=1.0)
    r = p.add_run(text.upper())
    set_font(r, size=10.5, bold=True, color="1F4D78")
    set_bottom_border(p)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    set_spacing(p, before=0, after=2.2, line=1.03)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=9.35, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, size=9.35)
    else:
        r = p.add_run(text)
        set_font(r, size=9.35)


def add_role(doc, company, dates, subtitle):
    p = doc.add_paragraph()
    set_spacing(p, before=3, after=0, line=1.0)
    r = p.add_run(company)
    set_font(r, size=10.2, bold=True)
    r = p.add_run(f" | {dates}")
    set_font(r, size=9.7, bold=True, color="555555")
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=2, line=1.0)
    r = p.add_run(subtitle)
    set_font(r, size=9.5, bold=True, color="1F4D78")


def add_inline_items(doc, items):
    p = doc.add_paragraph()
    set_spacing(p, before=0, after=3, line=1.04)
    for i, (label, value) in enumerate(items):
        if i:
            sep = p.add_run("  |  ")
            set_font(sep, size=9.15, color="666666")
        r = p.add_run(label)
        set_font(r, size=9.15, bold=True, color="1F4D78")
        r = p.add_run(value)
        set_font(r, size=9.15)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
styles["Normal"].font.size = Pt(10)
for style_name in ["List Bullet", "List Paragraph"]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(9.35)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(title, before=0, after=0, line=1.0)
r = title.add_run("BHARATH MYSORE")
set_font(r, size=17, bold=True, color="0B2545")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(contact, before=0, after=3, line=1.0)
r = contact.add_run("Bothell, WA | +1 (425) 503-6839 | bharathkumarmn@hotmail.com")
set_font(r, size=9.2, color="333333")

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(tag, before=0, after=6, line=1.0)
r = tag.add_run("Principal Systems Software Engineer | Storage Engines, Cloud Storage, and Hardware-Aware I/O")
set_font(r, size=10.2, bold=True, color="1F4D78")

add_heading(doc, "Targeted Profile")
p = doc.add_paragraph()
set_spacing(p, before=0, after=4, line=1.08)
r = p.add_run(
    "Principal-level storage and systems engineer with 20+ years across cloud storage platforms, C/C++ systems software, filesystems and disks, Linux kernel paths, RDMA, high-performance I/O, and hardware validation. Strong fit for Microsoft ESE modernization: built production storage capabilities spanning object, block, file, and key-value-oriented systems; led replication, snapshots/checkpoints, recovery, erasure coding, capacity efficiency, custom SSD/HDD filesystem work, fault containment, and storage-device integration. Brings the cross-layer judgment needed to reduce write amplification, improve endurance and capacity utilization, benchmark new storage architectures, and translate product-team requirements into durable platform capabilities."
)
set_font(r, size=9.65)

add_heading(doc, "ESE-Aligned Strengths")
add_inline_items(
    doc,
    [
        ("Storage engines: ", "log/recovery concepts, snapshots, checkpoints, write ordering, erasure coding, key-value-store POC"),
        ("Hardware-aware I/O: ", "QLC/SMR/Zoned-ready reasoning, SSD/HDD filesystems, NVMe/SPDK, io_uring, RDMA, iSCSI"),
    ],
)
add_inline_items(
    doc,
    [
        ("Efficiency: ", "write reduction, compression, tiering/archival, trimmed-block reclamation, capacity/TCO optimization"),
        ("Concurrency/perf: ", "C/C++, Folly fibers, multi-process serving, lock/contention-aware debugging, benchmarking"),
    ],
)
add_inline_items(
    doc,
    [
        ("Cloud platform: ", "OCI storage, Azure Storage POC, Kubernetes, Terraform, Prometheus, gRPC, Kafka/Flink/ZooKeeper"),
        ("Leadership: ", "cross-team architecture, prototyping, validation, production rollout, mentoring, partner alignment"),
    ],
)

add_heading(doc, "Professional Experience")
add_role(doc, "Oracle Corporation", "Jun 2016 - Present", "Cloud Storage Platforms - Unified Storage, Object Storage, and Block Storage")
oracle = [
    ("Unified storage architecture. ", "Unified storage architecture. Set architecture and execution direction for a converged storage platform serving object, block, and file workloads, including a replication layer shared across storage services."),
    ("Storage-engine data paths. ", "Storage-engine data paths. Designed and implemented production C++ erasure coding for Unified Storage Pool with quorum-based leader election and writes to a custom SSD/HDD filesystem."),
    ("Logging, checkpoint, and recovery fit. ", "Logging, checkpoint, and recovery fit. Built asynchronous cross-region replication for block volumes and volume groups using snapshots/checkpoints, write ordering, failover, and recovery semantics for disaster-recovery workloads."),
    ("Hardware-aware I/O modernization. ", "Hardware-aware I/O modernization. Built fully asynchronous erasure-coded and mirrored-replication paths with Folly fibers; applied SPDK/NVMe, io_uring, and Linux kernel I/O patterns in SSD/HDD-backed storage paths."),
    ("Write reduction and storage efficiency. ", "Write reduction and storage efficiency. Led cold-data compression reclaiming 20-30% of space, trimmed-block reclamation addressing 6-7% unreferenced data, and idle-volume optimization that produced U.S. Patent 11,412,043."),
    ("Scale and access-path performance. ", "Scale and access-path performance. Scaled backup systems, introduced multi-process iSCSI serving, and re-architected networking to reduce hops and support more attached volumes."),
    ("Reliability and validation. ", "Reliability and validation. Improved feature/volume/device fault containment, automated failed-disk and host-failure handling, drove recovery-time metrics, automated qualification, deployment safety, and touchless region buildout."),
    ("Benchmarking and experimentation. ", "Benchmarking and experimentation. Guided DP load-generation adoption, shaped customer-style LSM/disk/host/failure tests, tuned false positives, and used telemetry and performance baselines to prioritize engineering investments."),
    ("Storage technology pathfinding. ", "Storage technology pathfinding. Collaborated on a Binary Tree as a Service key-value-store proof of concept and evaluated existing-system performance to guide design; participated in Deep Archive vendor discussions and tape-backed cold-storage tradeoffs."),
    ("Cross-team platform execution. ", "Cross-team platform execution. Worked across Block Storage, Volume Service, USP Client, Extent Server, and Object Storage teams to integrate Block workloads with the unified platform and turn ambiguous requirements into designs, reviews, tests, and bring-up plans."),
]
for prefix, text in oracle:
    add_bullet(doc, text, prefix)

doc.add_page_break()
add_role(doc, "Microsoft Corporation", "Sep 2005 - May 2016", "Azure Storage, Windows Filesystems/Disks, Visual Studio Parallel/GPU Debugging, HoloLens/Xbox")
ms = [
    ("Azure Storage RDMA. ", "Azure Storage RDMA. Designed and implemented an RDMA library for a proof-of-concept Azure Storage backend, connecting low-latency transport behavior with cloud storage architecture."),
    ("Filesystem and disk internals. ", "Filesystem and disk internals. Served as Windows OS, filesystem, and disk subject-matter expert, leading kernel-dump analysis, root-cause investigation, reverse engineering, code reviews, and customer-facing resolution of disk and cluster failures."),
    ("Concurrency and performance tooling. ", "Concurrency and performance tooling. Led Visual Studio Cluster Debugger for HPC/MPI, Parallel Debugger for concurrency, and GPU Debugger for CUDA/OpenGL/DirectX; built engine and UI test libraries and coordinated vendor test migration."),
    ("Hardware/software delivery. ", "Hardware/software delivery. Owned HoloLens depth-camera functionality from prototype through public release; developed hardware-verification drivers/tools for Surface Hub and Kinect and graphics-analysis prototype work for Xbox."),
]
for prefix, text in ms:
    add_bullet(doc, text, prefix)

add_role(doc, "Intel Technology India Private Limited", "Mar 2002 - Sep 2005", "CPU, Cache, Memory, Platform Validation, and Linux Kernel/Driver Development")
intel = [
    ("CPU/cache/memory validation. ", "CPU/cache/memory validation. Designed validation tools for CPU cores and cache/memory subsystems, including Linux kernel programming and drivers for test cards and chipsets."),
    ("Hardware/OS root cause. ", "Hardware/OS root cause. Drove complex hardware and operating-system defects to closure through debugging and targeted patches; developed network-card management/configuration software for Windows and Linux."),
]
for prefix, text in intel:
    add_bullet(doc, text, prefix)

add_role(doc, "Infineon Technologies (formerly Siemens Semiconductor)", "Mar 2000 - Mar 2002", "Systems Software Engineer")
add_bullet(doc, "Embedded control systems. Developed software to configure, monitor, and control telephone-exchange cards at the hardware and telecommunications interface.", "Embedded control systems. ")

add_heading(doc, "Selected Fit for Microsoft ESE")
fit = [
    ("Modern storage media. ", "Modern storage media. Direct SSD/HDD, NVMe/SPDK, io_uring, RDMA, Linux kernel, filesystem, disk, and device-driver background for evaluating QLC, SMR, Zoned Storage, endurance, and predictable performance tradeoffs."),
    ("Storage engine concepts. ", "Storage engine concepts. Practical experience with replication logs, snapshots/checkpoints, recovery semantics, fault isolation, custom filesystem integration, key-value-store exploration, and LSM/disk/host/failure validation."),
    ("Platform leadership. ", "Platform leadership. Proven record driving architecture across storage, client, metadata, service, and operations teams while mentoring engineers and translating product requirements into platform capabilities."),
]
for prefix, text in fit:
    add_bullet(doc, text, prefix)

add_heading(doc, "Education and Patent")
p = doc.add_paragraph()
set_spacing(p, before=0, after=0, line=1.0)
r = p.add_run("B.E., Computer Science and Engineering, Mangalore University")
set_font(r, size=9.4)
r = p.add_run("  |  U.S. Patent 11,412,043 - idle-volume backend capacity optimization")
set_font(r, size=9.4, color="333333")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_spacing(footer, before=0, after=0, line=1.0)
r = footer.add_run("Tailored for Microsoft Extensible Storage Engine (ESE)")
set_font(r, size=8.2, color="666666")

doc.save(OUT)
print(OUT)
