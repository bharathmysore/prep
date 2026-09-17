from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

ROOT = Path('/Users/bmysoren/prep')
doc = Document(ROOT / 'resume/Bharath_Mysore_OpenAI_Compute_Storage.docx')
for child in list(doc.element.body):
    if child.tag != qn('w:sectPr'):
        doc.element.body.remove(child)

def p(text, style=None):
    para = doc.add_paragraph(text, style)
    para.paragraph_format.widow_control = True
    return para

def bullet(text):
    para = p(text, 'List Bullet')
    para.paragraph_format.left_indent = Inches(0.15)
    para.paragraph_format.first_line_indent = Inches(-0.15)
    para.paragraph_format.keep_together = True

def heading(text):
    para = p(text, 'Heading 1')
    para.paragraph_format.space_before = Pt(10)
    para.paragraph_format.space_after = Pt(5)

def role(company, dates, scope):
    para = p('', 'Heading 2')
    para.add_run(company).bold = True
    para.add_run(' | ' + dates).bold = False
    para = p(scope)
    para.runs[0].italic = True
    para.paragraph_format.keep_with_next = True

p('BHARATH MYSORE', 'Title')
p('Bothell, WA | +1 (425) 503-6839 | bharathkumarmn@hotmail.com')
para = p('DISTRIBUTED SYSTEMS AND HIGH PERFORMANCE COMPUTING TOOLS')
para.runs[0].bold = True
p('Systems software engineer and technical leader with 20+ years across cloud infrastructure, distributed storage, HPC and GPU developer tools, and hardware platforms. Combine hands-on C/C++ development with cross-team architecture and delivery from prototype to release. Experience includes Oracle cloud storage, an Azure Storage RDMA prototype, Visual Studio cluster and parallel debugger test teams, and end-to-end HoloLens device integration.')
heading('Technical Strengths')
p('Distributed storage and replication; concurrent and parallel programming; HPC/MPI and GPU debugging tools; RDMA; Linux kernel and device drivers; systems root-cause analysis; qualification and deployment automation; engineering mentorship.')
heading('Professional Experience')
role('Oracle Corporation', 'Jun 2016 - Present', 'Oracle Cloud Infrastructure | Unified Storage and Block Storage')
for text in [
    'Design and drive cross-team initiatives for a unified storage platform serving object, block, and file workloads, including shared replication and erasure coding for object storage.',
    'Build block-storage support on the unified stack and drive migration from the existing backend; contribute to deep archival capabilities. Align work across teams on a new storage platform.',
    'Lead cross-region and asynchronous volume-group replication from design through delivery for disaster recovery. Ownership spans replication, leader election, snapshots, and backup.',
    'Scale the backup backend with service growth and implement application-controlled snapshotting. Introduce multiprocessing for iSCSI serving and rearchitect networking to reduce hops and support more attached volumes.',
    'Increase sellable capacity through cold-data compression, reclaiming 20-30% of storage space. Optimize idle-volume backend capacity through work associated with U.S. Patent 11,412,043; implement trimmed-block reclamation.',
    'Refactor the backend to contain failures within a feature, volume, or device. Automate failed-disk handling and detection and handling of single-host failures.',
    'Drive automated qualification and deployments through CI/CD, together with single-click region buildout to reduce manual infrastructure delivery work.',
    'Design and develop a near-real-time marketplace metering service from prototype to public release; help redirect and implement a data-processing service through internal release.',
]:
    bullet(text)

doc.add_page_break()
heading('Professional Experience Continued')
role('Microsoft Corporation', 'Sep 2005 - May 2016', 'Visual Studio | Azure Storage | Windows Systems | Device Platforms')
for text in [
    'Led Visual Studio Cluster, Parallel, and GPU debugger test teams spanning HPC/MPI, concurrent programming, CUDA, OpenGL, and DirectX. Built local and remote teams, contributed to hiring, and mentored engineers.',
    'Implemented engine test libraries for cluster and GPU debuggers and architected a reusable debugger UI test library. Worked with vendors to migrate tests across product cycles.',
    'Designed and implemented an RDMA library for a proof-of-concept Azure Storage backend, applying systems programming expertise to storage transport.',
    'Owned end-to-end HoloLens depth-camera functionality from prototype through public announcement. Developed hardware-validation drivers and tools for Surface Hub and Kinect.',
    'Worked on Xbox 360 graphics emulation for Xbox One and a frame-capture tool for collecting graphics metrics under controlled parameters.',
    'Served as a Windows OS, filesystem, and disk subject-matter expert. Used kernel-dump analysis, reverse engineering, and code review to diagnose disk and cluster failures with customers and development teams.',
]:
    bullet(text)
role('Intel Technology India Private Limited', 'Mar 2002 - Sep 2005', 'Systems Software and Hardware Validation')
for text in [
    'Designed and developed CPU-core and cache/memory-subsystem validation tools, including Linux kernel programming and drivers for test cards and chipsets.',
    'Debugged hardware and operating-system defects, drove issues to closure, and wrote targeted patches. Developed network-card management and configuration software for Windows and Linux.',
]:
    bullet(text)
heading('Earlier Experience')
p('Infineon Technologies (formerly Siemens Semiconductor) | Mar 2000 - Mar 2002. Developed systems software to configure, monitor, and control telephone-exchange cards.')
heading('Technical Skills')
for label, value in [
    ('Languages', 'C/C++; Java, Python, and C# as needed.'),
    ('Infrastructure', 'Kubernetes, Docker, Terraform, Helm, Prometheus, gRPC, Folly, Boost, ZooKeeper, Kafka, and Flink.'),
    ('Systems', 'Object, block, and file storage; replication; erasure coding; snapshots; disaster recovery; iSCSI; RDMA; Linux and Windows systems debugging.'),
]:
    para = p('')
    para.add_run(label + ': ').bold = True
    para.add_run(value)
heading('Education')
p('B.E., Computer Science and Engineering, Mangalore University')
doc.core_properties.title = 'Bharath Mysore Resume for Microsoft Research Principal Software Engineer'
doc.core_properties.subject = 'Microsoft job 200052038 | Research infrastructure and distributed systems'
doc.core_properties.author = 'Bharath Mysore'
out = ROOT / 'resume/Bharath_Mysore_Microsoft_Research_Principal_Software_Engineer_Resume.docx'
doc.save(out)
print(out)
