from pathlib import Path
from docx import Document
from docx.shared import Pt

ROOT = Path('/Users/bmysoren/prep')
doc = Document(ROOT / 'resume/Bharath_Mysore_OpenAI_Compute_Storage.docx')
replacements = {
    'STORAGE INFRASTRUCTURE AND DISTRIBUTED SYSTEMS':
        'INFRASTRUCTURE RELIABILITY AND PLATFORM ENGINEERING',
    'Systems software engineer and technical leader with 20+ years of experience across cloud storage, operating systems, and developer tools. At Oracle Cloud Infrastructure, design and implement storage capabilities spanning object, block, and file workloads, shared replication, cross-region disaster recovery, and capacity efficiency. Combine hands-on C/C++ development with cross-team delivery, deep systems debugging, and automation of storage operations.':
        'Infrastructure engineer and technical leader with 20+ years of experience across hyperscale cloud storage, distributed systems, and hardware-to-software platforms. At Oracle Cloud Infrastructure, drive cross-team work in reliability automation, capacity efficiency, disaster recovery, and region buildout. Combine hands-on systems development and root-cause analysis with technical direction, engineering mentorship, and delivery of ambiguous programs from prototype to release.',
    'Oracle Cloud Infrastructure | Unified Storage and Block Storage':
        'Oracle Cloud Infrastructure | Storage Reliability and Platform Engineering',
    'Design and drive cross-team initiatives for a unified storage platform serving object, block, and file workloads, including a shared replication layer and erasure coding for object storage.':
        'Drive automated qualification and deployment through CI/CD and single-click region buildout, turning repeated infrastructure delivery tasks into automated workflows.',
    'Build block-storage support on the unified stack and drive migration from the existing backend; contribute to deep archival capabilities for long-term storage.':
        'Automate failed-disk handling and detection and handling of single-host failures. Refactor the backend to contain failures within a feature, volume, or device and limit their operational impact.',
    'Develop cross-region replication for disaster recovery, including asynchronous replication of volume groups. Ownership spans replication and leader election, snapshots, backup, and cross-region replication.':
        'Lead capacity-efficiency initiatives: cold-data compression reclaimed 20-30% of storage space, increasing sellable capacity. Optimize idle-volume backend capacity through work associated with U.S. Patent 11,412,043.',
    'Implement application-controlled snapshotting and scale the backup backend with service growth, supporting data protection across the storage lifecycle.':
        'Design and drive cross-team initiatives for a unified platform serving object, block, and file workloads, including shared replication and erasure coding for object storage.',
    'Increase sellable capacity through cold-data compression, reclaiming 20-30% of storage space. Optimize idle-volume backend capacity through work associated with U.S. Patent 11,412,043.':
        'Lead cross-region replication and asynchronous volume-group replication from design through delivery, strengthening disaster-recovery capabilities. Ownership includes replication, leader election, snapshots, and backup.',
    'Implement reclamation of trimmed user blocks to recover unreferenced storage capacity.':
        'Scale the backup backend with service growth and implement application-controlled snapshotting. Build block-storage support on the unified stack and drive migration from the existing backend.',
    'Scale iSCSI serving through multiprocessing and rearchitect the networking subsystem to reduce hops and increase the number of attached volumes.':
        'Scale iSCSI serving through multiprocessing and rearchitect networking to reduce hops and increase the number of attached volumes, addressing backend and access-path constraints.',
    'Refactor the storage backend to contain failures within a feature, volume, or device. Automate failed-disk handling and detection and handling of single-host failures.':
        'Implement trimmed-block reclamation to recover unreferenced storage capacity and contribute to deep archival capabilities for long-term retention.',
    'Drive automated qualification and deployment through CI/CD, together with single-click region buildout to reduce manual operational work.':
        'Design and develop a near-real-time marketplace metering service from prototype to public release; help redirect and implement a data-processing service through internal release.',
    'Azure Storage | Visual Studio | Windows Systems | Device Platforms':
        'Windows Reliability | Azure Storage | Developer Tools | Device Platforms',
    'Designed and implemented an RDMA library for a proof-of-concept Azure Storage backend, connecting systems programming experience with storage transport performance.':
        'Served as a Windows operating-system, filesystem, and disk subject-matter expert. Led kernel-dump analysis and root-cause investigations, working with customers and development teams to resolve disk and cluster failures.',
    'Led Visual Studio Cluster Debugger for HPC/MPI, Parallel Debugger, and GPU Debugger projects spanning CUDA, OpenGL, and DirectX. Built teams of five local and ten remote engineers; contributed to hiring and mentoring.':
        'Led Visual Studio Cluster, Parallel, and GPU debugger test teams across HPC/MPI, concurrency, CUDA, OpenGL, and DirectX. Built local and remote teams, contributed to hiring, and mentored engineers.',
    'Implemented engine test libraries for cluster and GPU debuggers and architected a reusable debugger UI test library. Coordinated with vendors to migrate tests across product cycles.':
        'Implemented reusable engine test libraries for cluster and GPU debuggers and architected a debugger UI test library. Coordinated with vendors to migrate validation across product cycles.',
    'Diagnosed Windows operating-system, filesystem, disk, and cluster failures through kernel-dump analysis, reverse engineering, and code review. Worked with development teams and customers to resolve root causes.':
        'Designed and implemented an RDMA library for a proof-of-concept Azure Storage backend, applying systems programming expertise to storage transport.',
    'Owned end-to-end HoloLens depth-camera functionality from prototype through public announcement. Developed hardware-validation drivers and tools for Surface Hub and Kinect, and worked on Xbox graphics emulation and frame-capture tooling.':
        'Owned HoloLens depth-camera functionality end to end from prototype through public announcement. Developed hardware-validation drivers and tools for Surface Hub and Kinect.',
    'Additional Cloud Service Delivery': 'Earlier Experience',
    'At Oracle, designed and developed a near-real-time marketplace software-metering service from prototype to public release. Helped redirect and implement a data-processing service, taking it from prototype through internal release.':
        'Infineon Technologies (formerly Siemens Semiconductor) | Mar 2000 - Mar 2002. Developed systems software to configure, monitor, and control telephone-exchange cards.',
}
seen = set()
for para in doc.paragraphs:
    if para.text in replacements:
        old = para.text
        if len(para.runs) == 1:
            para.runs[0].text = replacements[old]
        else:
            para.text = replacements[old]
        seen.add(old)
assert seen == set(replacements), set(replacements) - seen

for para in doc.paragraphs:
    if para.text.startswith('Storage and systems:'):
        para.clear()
        para.add_run('Reliability and systems: ').bold = True
        para.add_run('Fault containment; disk and host failure automation; disaster recovery; replication; snapshots and backup; capacity efficiency; Linux kernel and device drivers; RDMA; concurrent and parallel programming.')
    elif para.text.startswith('Infrastructure and tooling:'):
        para.clear()
        para.add_run('Infrastructure and tooling: ').bold = True
        para.add_run('Kubernetes, Terraform, Docker, Helm, Prometheus, CI/CD, gRPC, Folly, Boost, ZooKeeper, Kafka, and Flink.')

doc.add_paragraph('Education', 'Heading 1')
doc.add_paragraph('B.E., Computer Science and Engineering, Mangalore University')
doc.core_properties.title = 'Bharath Mysore Resume for Cerebras Principal SRE AI Inference'
doc.core_properties.subject = 'Infrastructure reliability and platform engineering'
out = ROOT / 'resume/Bharath_Mysore_Cerebras_Principal_SRE_AI_Inference_Resume.docx'
doc.save(out)
print(out)
