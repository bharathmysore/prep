from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/bmysoren/prep/resume/Bharath_Mysore_OpenAI_Compute_Storage.docx"


def set_font(run, name="Calibri", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_p_spacing(paragraph, before=0, after=3, line=1.06):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_bottom_border(paragraph, color="D9E2F3", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_heading(doc, text):
    p = doc.add_paragraph()
    set_p_spacing(p, before=5, after=3, line=1.0)
    r = p.add_run(text.upper())
    set_font(r, size=10.5, bold=True, color="1F4D78")
    set_bottom_border(p)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    set_p_spacing(p, before=0, after=2.3, line=1.03)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=9.4, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, size=9.4)
    else:
        r = p.add_run(text)
        set_font(r, size=9.4)
    return p


def add_role(doc, company, dates, subtitle):
    p = doc.add_paragraph()
    set_p_spacing(p, before=3, after=0, line=1.0)
    left = p.add_run(company)
    set_font(left, size=10.2, bold=True, color="000000")
    right = p.add_run(f" | {dates}")
    set_font(right, size=9.7, bold=True, color="555555")
    p2 = doc.add_paragraph()
    set_p_spacing(p2, before=0, after=2, line=1.0)
    r = p2.add_run(subtitle)
    set_font(r, size=9.5, bold=True, color="1F4D78")


def add_inline_items(doc, items):
    p = doc.add_paragraph()
    set_p_spacing(p, before=0, after=3, line=1.04)
    for i, (label, value) in enumerate(items):
        if i:
            sep = p.add_run("  |  ")
            set_font(sep, size=9.2, color="666666")
        r = p.add_run(label)
        set_font(r, size=9.2, bold=True, color="1F4D78")
        r = p.add_run(value)
        set_font(r, size=9.2)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.62)
section.right_margin = Inches(0.62)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
styles["Normal"].font.size = Pt(10)

for style_name in ["List Bullet", "List Paragraph"]:
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(9.4)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_p_spacing(title, before=0, after=0, line=1.0)
r = title.add_run("BHARATH MYSORE")
set_font(r, size=17, bold=True, color="0B2545")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_p_spacing(contact, before=0, after=3, line=1.0)
r = contact.add_run("Bothell, WA | +1 (425) 503-6839 | bharathkumarmn@hotmail.com")
set_font(r, size=9.2, color="333333")

tag = doc.add_paragraph()
tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_p_spacing(tag, before=0, after=6, line=1.0)
r = tag.add_run("Senior Storage Infrastructure Engineer | Distributed Storage, Replication, and Production Reliability")
set_font(r, size=10.2, bold=True, color="1F4D78")

add_heading(doc, "Targeted Profile")
p = doc.add_paragraph()
set_p_spacing(p, before=0, after=4, line=1.08)
r = p.add_run(
    "Storage and systems software engineer with 20+ years building production infrastructure across cloud object, block, and file storage; cross-region data movement; lifecycle and capacity management; high-performance I/O; Linux/kernel paths; and device integration. At Oracle Cloud Infrastructure, set architecture and execution direction for a unified storage platform and shared replication layer, with hands-on C++ delivery across erasure coding, custom SSD/HDD filesystems, snapshots/checkpoints, failover, recovery, and automated operations. Strong fit for OpenAI Compute Storage work spanning object stores, federation layers, Kubernetes-based services, durability, availability, and end-to-end production ownership."
)
set_font(r, size=9.7)

add_heading(doc, "Role-Aligned Capabilities")
add_inline_items(
    doc,
    [
        ("Storage: ", "object, block, file, key-value concepts, erasure coding, replication, snapshots, checkpoints, archival"),
        ("Data movement: ", "cross-region replication, write ordering, failover, recovery, disaster recovery"),
    ],
)
add_inline_items(
    doc,
    [
        ("Systems code: ", "production C/C++, Folly fibers, SPDK, io_uring, Linux kernel I/O, RDMA, device drivers"),
        ("Platform ops: ", "Kubernetes, Terraform, Docker, Helm, Prometheus, gRPC, ZooKeeper, Kafka, Flink"),
    ],
)
add_inline_items(
    doc,
    [
        ("Reliability: ", "fault containment, disk/host-failure automation, corruption/recovery thinking, qualification gates"),
        ("Leadership: ", "cross-team architecture, ambiguous V1 systems, production release, mentoring and team formation"),
    ],
)

add_heading(doc, "Professional Experience")
add_role(doc, "Oracle Corporation", "Jun 2016 - Present", "Cloud Storage Platforms - Unified Storage, Object Storage, and Block Storage")
oracle_bullets = [
    ("Unified storage architecture. ", "Unified storage architecture. Set architecture and execution direction for storage capabilities serving object, block, and file workloads, including a replication layer shared across storage services."),
    ("Object-store data path. ", "Object-store data path. Designed and implemented production C++ erasure coding for Unified Storage Pool with quorum-based leader election and writes to a custom SSD/HDD filesystem."),
    ("Cross-region data movement. ", "Cross-region data movement. Built asynchronous replication for block volumes and volume groups with snapshots/checkpoints, write ordering, failover, and recovery across OCI Block Storage."),
    ("Backend unification and lifecycle. ", "Backend unification and lifecycle. Drove block-storage bring-up and migration to a unified stack and delivered deep archival block storage, broadening retention and cost-management options."),
    ("High-performance storage I/O. ", "High-performance storage I/O. Built fully asynchronous erasure-coded and mirrored-replication paths with Folly fibers; used SPDK and io_uring for SSD/NVMe integration and Linux kernel I/O for HDD-backed object storage."),
    ("Scale and access-path efficiency. ", "Scale and access-path efficiency. Scaled the backup backend, introduced multi-process iSCSI serving, and re-architected networking to reduce hops and support more attached volumes."),
    ("Durability and operations. ", "Durability and operations. Improved fault containment and automated failed-disk and host-failure handling; drove automated qualification, deployment, and single-click region buildout."),
    ("Storage efficiency. ", "Storage efficiency. Cold-data compression reclaimed 20-30% of space; trimmed-block reclamation addressed 6-7% of unreferenced data; idle-volume optimization produced U.S. Patent 11,412,043."),
    ("Storage service exploration. ", "Storage service exploration. Partnered on a Binary Tree as a Service key-value-store proof of concept and evaluated existing performance baselines to guide design direction."),
]
for prefix, text in oracle_bullets:
    add_bullet(doc, text, prefix)

doc.add_page_break()
add_role(doc, "Microsoft Corporation", "Sep 2005 - May 2016", "Azure Storage, Systems, Developer Tools, Windows, and Hardware-to-Software Platforms")
ms_bullets = [
    ("Azure Storage and RDMA. ", "Azure Storage and RDMA. Designed and implemented an RDMA library for proof-of-concept Azure Storage backend infrastructure."),
    ("HPC and GPU developer tools. ", "HPC and GPU developer tools. Led Visual Studio Cluster (HPC/MPI), Parallel, and GPU debugger projects; built engine and UI test libraries, led local and remote teams, and supported CUDA, OpenGL, and DirectX debugging."),
    ("Cross-layer systems debugging. ", "Cross-layer systems debugging. Served as a Windows OS, filesystem, and disk subject-matter expert, leading kernel-dump analysis, root-cause investigation, reverse engineering, code reviews, and customer issue resolution for disk and cluster failures."),
    ("Hardware-to-software delivery. ", "Hardware-to-software delivery. Owned HoloLens depth-camera functionality from prototype through public release; developed validation drivers and tools for Surface Hub and Kinect and a graphics-emulation proof of concept for Xbox."),
]
for prefix, text in ms_bullets:
    add_bullet(doc, text, prefix)

add_role(doc, "Intel Technology India Private Limited", "Mar 2002 - Sep 2005", "Systems Software and Hardware Validation")
intel_bullets = [
    ("CPU and platform validation. ", "CPU and platform validation. Designed validation tools for CPU cores and cache/memory subsystems, including Linux kernel programming and drivers for test cards and chipsets."),
    ("Hardware/OS root cause. ", "Hardware/OS root cause. Drove complex hardware and operating-system defects to closure through debugging and targeted patches; built network-card management software for Windows and Linux."),
]
for prefix, text in intel_bullets:
    add_bullet(doc, text, prefix)

add_role(doc, "Infineon Technologies (formerly Siemens Semiconductor)", "Mar 2000 - Mar 2002", "Systems Software Engineer")
add_bullet(doc, "Embedded control systems. Developed software to configure, monitor, and control telephone-exchange cards at the hardware and telecommunications interface.", "Embedded control systems. ")

add_heading(doc, "Leadership and Operating Model")
leadership = [
    ("Technical direction. ", "Technical direction. Sets architecture and execution direction for cross-team platform work and drives ambiguous initiatives from proof of concept through production release."),
    ("Operational ownership. ", "Operational ownership. Leads complex root-cause analysis, code reviews, and customer-facing resolution while turning recurring failures into qualification, deployment, and remediation automation."),
    ("Team development. ", "Team development. Built and led local and remote engineering teams, hired and mentored engineers, and created reusable engine and UI test libraries."),
]
for prefix, text in leadership:
    add_bullet(doc, text, prefix)

add_heading(doc, "Additional Delivery")
add_bullet(doc, "Data-intensive cloud services. Built a near-real-time marketplace metering service from prototype to public release and helped take a data-processing service from prototype to internal release.", "Data-intensive cloud services. ")

add_heading(doc, "Education and Patent")
edu = doc.add_paragraph()
set_p_spacing(edu, before=0, after=0, line=1.0)
r = edu.add_run("B.E., Computer Science and Engineering, Mangalore University")
set_font(r, size=9.4)
r = edu.add_run("  |  U.S. Patent 11,412,043 - idle-volume backend capacity optimization")
set_font(r, size=9.4, color="333333")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_p_spacing(footer, before=0, after=0, line=1.0)
r = footer.add_run("Tailored for OpenAI Software Engineer, Compute - Storage")
set_font(r, size=8.2, color="666666")

doc.save(OUT)
print(OUT)
